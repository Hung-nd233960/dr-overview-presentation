# -*- coding: utf-8 -*-
import re
from docx import Document
from docx.shared import Pt, Cm
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from script_data import SLIDES, HANDOVERS, RUNSHEET_BJ, RUNSHEET_WU

GLOSS = [
("digital radiography (DR)", "chụp X-quang kỹ thuật số"), ("computed radiography (CR)", "X-quang điện toán (CR)"),
("film-screen radiography", "chụp X-quang phim – màn tăng quang"), ("imaging plate / photostimulable phosphor", "tấm ghi ảnh / chất phosphor kích thích quang"),
("flat-panel detector (indirect / direct)", "đầu thu phẳng (gián tiếp / trực tiếp)"), ("scintillator", "chất nhấp nháy"),
("photoconductor (amorphous selenium)", "chất quang dẫn (selen vô định hình)"), ("thin-film transistor (TFT) array", "ma trận transistor màng mỏng (TFT)"),
("photomultiplier tube (PMT)", "ống nhân quang (PMT)"), ("analog-to-digital converter (ADC)", "bộ chuyển đổi tương tự – số (ADC)"),
("optical density (OD)", "mật độ quang học (OD)"), ("characteristic (H&D) curve", "đường cong đặc trưng (H&D)"),
("toe / straight-line portion / shoulder", "vùng chân / vùng tuyến tính / vùng vai"), ("exposure latitude / dynamic range", "độ rộng phơi sáng / dải động"),
("tube voltage (kVp) / tube current-time (mAs)", "điện áp đỉnh (kVp) / tích dòng – thời gian (mAs)"), ("automatic exposure control (AEC)", "kiểm soát phơi sáng tự động (AEC)"),
("quantum mottle (noise)", "nhiễu lượng tử (hạt nhiễu)"), ("spatial resolution / contrast resolution", "độ phân giải không gian / độ phân giải tương phản"),
("detective quantum efficiency (DQE)", "hiệu suất lượng tử phát hiện (DQE)"), ("exposure index / deviation index", "chỉ số phơi sáng (EI) / chỉ số sai lệch (DI)"),
("dose creep", "hiện tượng tăng liều dần (dose creep)"), ("digital fluoroscopy / digital subtraction angiography", "chiếu huỳnh quang số / chụp mạch số hóa xóa nền (DSA)"),
("digital mammography / computer-aided detection", "chụp nhũ ảnh số / phát hiện có hỗ trợ máy tính (CAD)"),
("PACS / RIS / HIS", "hệ thống lưu trữ và truyền ảnh (PACS) / hệ thống thông tin chẩn đoán hình ảnh (RIS) / hệ thống thông tin bệnh viện (HIS)"),
("lossless / lossy compression", "nén không mất dữ liệu / nén có mất dữ liệu"), ("quality assurance / quality control", "đảm bảo chất lượng / kiểm soát chất lượng"),
("ALARA", "thấp nhất có thể đạt được một cách hợp lý (ALARA)"),
]


def shade(c, fill):
    tcPr = c._element.get_or_add_tcPr()
    sh = OxmlElement('w:shd')
    sh.set(qn('w:val'), 'clear'); sh.set(qn('w:color'), 'auto'); sh.set(qn('w:fill'), fill)
    tcPr.append(sh)


def setcell(c, text, bold=False, size=11, fill=None):
    c.text = ''
    r = c.paragraphs[0].add_run(text)
    r.bold = bold; r.font.size = Pt(size); r.font.name = 'Calibri'
    if fill: shade(c, fill)


def newdoc():
    d = Document()
    for s in d.sections:
        s.left_margin = s.right_margin = Cm(1.8); s.top_margin = s.bottom_margin = Cm(1.6)
    st = d.styles['Normal']; st.font.name = 'Calibri'; st.font.size = Pt(11)
    return d


def build_script():
    d = newdoc(); md = []
    d.add_heading('Script (EN | VI) — Digital Radiography: An Overview', 0)
    md.append('# Script (EN | VI) — Digital Radiography: An Overview\n')
    intro = 'Target 17:00 including two activities. Speakers: A, B, C, D (replace with names). Lines in [brackets] are stage directions. The Vietnamese is a natural spoken translation, not word for word; please have it read by a native speaker.'
    d.add_paragraph(intro); md.append(intro + '\n')
    for num, title, spk, tm, rows in SLIDES:
        label = f'Slide {num}' if isinstance(num, int) else str(num)
        d.add_heading(f'{label}: {title}', 2)
        d.add_paragraph(f'Speaker: {spk}   ·   Time: {tm}')
        t = d.add_table(rows=1, cols=2); t.style = 'Table Grid'
        setcell(t.rows[0].cells[0], 'English (say aloud)', True, 11, 'D9E7F5'); setcell(t.rows[0].cells[1], 'Tiếng Việt (meaning)', True, 11, 'D9E7F5')
        md.append(f'## {label}: {title}\n\n**Speaker:** {spk} · **Time:** {tm}\n\n| English (say aloud) | Tiếng Việt (meaning) |\n|---|---|')
        for en, vi in rows:
            r = t.add_row().cells; setcell(r[0], en); setcell(r[1], vi)
            md.append(f'| {en} | {vi} |')
        md.append('')
    d.add_heading('Handover lines', 1); md.append('## Handover lines\n\n| From → to | When | Line |\n|---|---|---|')
    t = d.add_table(rows=1, cols=3); t.style = 'Table Grid'
    for i, x in enumerate(['From → to', 'When', 'Line']): setcell(t.rows[0].cells[i], x, True, 11, 'D9E7F5')
    for a, b, c in HANDOVERS:
        r = t.add_row().cells; setcell(r[0], a); setcell(r[1], b); setcell(r[2], c); md.append(f'| {a} | {b} | {c} |')
    for title, steps in (('Run sheet — Activity 1: Radiologist Eye Test', RUNSHEET_WU), ('Run sheet — Activity 2: X-ray Blackjack', RUNSHEET_BJ)):
        d.add_heading(title, 1); md.append(f'\n## {title}\n')
        for i, s in enumerate(steps, 1):
            d.add_paragraph(s, style='List Number'); md.append(f'{i}. {s}')
    d.add_heading('Glossary EN | VI', 1); md.append('\n## Glossary EN | VI\n\n| English | Tiếng Việt |\n|---|---|')
    t = d.add_table(rows=1, cols=2); t.style = 'Table Grid'
    setcell(t.rows[0].cells[0], 'English', True, 11, 'D9E7F5'); setcell(t.rows[0].cells[1], 'Tiếng Việt', True, 11, 'D9E7F5')
    for en, vi in GLOSS:
        r = t.add_row().cells; setcell(r[0], en); setcell(r[1], vi); md.append(f'| {en} | {vi} |')
    d.save('../Script_EN_VI.docx')
    open('../Script_EN_VI.md', 'w', encoding='utf-8').write('\n'.join(md) + '\n')


def inline(p, text, size=11):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\*[^*]+\*)', text):
        if not part: continue
        if part.startswith('**'): r = p.add_run(part[2:-2]); r.bold = True
        elif part.startswith('`'): r = p.add_run(part[1:-1]); r.font.name = 'Courier New'
        elif part.startswith('*') and len(part) > 2: r = p.add_run(part[1:-1]); r.italic = True
        else: r = p.add_run(part)
        r.font.size = Pt(size)


def md_to_docx(src, dst):
    d = newdoc(); lines = open(src, encoding='utf-8').read().split('\n'); i = 0
    while i < len(lines):
        l = lines[i]
        if l.startswith('# '): d.add_heading(l[2:], 0)
        elif l.startswith('## '): d.add_heading(l[3:], 1)
        elif l.startswith('### '): d.add_heading(l[4:], 2)
        elif l.startswith('|'):
            rows = []
            while i < len(lines) and lines[i].startswith('|'):
                if not re.match(r'^\|[-| ]+\|$', lines[i]): rows.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            n = len(rows[0]); t = d.add_table(rows=len(rows), cols=n); t.style = 'Table Grid'
            for ri, row in enumerate(rows):
                for ci in range(n):
                    c = t.rows[ri].cells[ci]; c.text = ''
                    inline(c.paragraphs[0], row[ci] if ci < len(row) else '', 10)
                    if ri == 0:
                        for r in c.paragraphs[0].runs: r.bold = True
                        shade(c, 'D9E7F5')
            d.add_paragraph(); continue
        elif re.match(r'^\s*- ', l):
            indent = len(l) - len(l.lstrip()); p = d.add_paragraph(style='List Bullet 2' if indent >= 2 else 'List Bullet'); inline(p, l.strip()[2:])
        elif re.match(r'^\d+\. ', l):
            p = d.add_paragraph(style='List Number'); inline(p, re.sub(r'^\d+\. ', '', l))
        elif l.strip() == '': pass
        else:
            p = d.add_paragraph(); inline(p, l)
        i += 1
    d.save(dst)


build_script()
md_to_docx('../Knowledge_Guide.md', '../Knowledge_Guide.docx')
print('docs ok')
